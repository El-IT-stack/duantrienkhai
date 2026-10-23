import docx

doc = docx.Document('Bao_Cao_Trien_Khai_Quan_Tri_He_Thong_Quan_Ly_Du_An (1).docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")
for i, p in enumerate(doc.paragraphs):
    blips = p._element.xpath('.//a:blip')
    if blips:
        for b in blips:
            embed_id = b.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
            target = doc.part.rels[embed_id].target_ref if embed_id in doc.part.rels else 'unknown'
            print(f"Paragraph {i}: Image target={target}, embed={embed_id}, text='{p.text}'")

for ti, table in enumerate(doc.tables):
    for ri, r in enumerate(table.rows):
        for ci, c in enumerate(r.cells):
            blips = c._element.xpath('.//a:blip')
            if blips:
                for b in blips:
                    embed_id = b.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                    target = doc.part.rels[embed_id].target_ref if embed_id in doc.part.rels else 'unknown'
                    print(f"Table {ti} R{ri}C{ci}: Image target={target}, text='{c.text[:40]}'")
